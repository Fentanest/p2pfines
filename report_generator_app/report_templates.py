from __future__ import annotations

from datetime import datetime
from typing import Any
import json
import re

import openpyxl
from openpyxl.utils import range_boundaries

from .schema import SOURCE_SPECS, WATCH_REPORT_SHEETS, WORK_REPORT_SHEETS, default_columns_for_target
from .sheets_client import GoogleSheetsClient
from .paths import RESOURCE_ROOT


WORKBOOK_PATH = RESOURCE_ROOT / "업무보고서 작성기_v5.4_통합.xlsx"
WORK_APPS_SCRIPT_PATH = RESOURCE_ROOT / "apps_script" / "Code_work.gs"
WATCH_APPS_SCRIPT_PATH = RESOURCE_ROOT / "apps_script" / "Code_watch.gs"


def setup_online_documents(
    client: GoogleSheetsClient,
    work_spreadsheet_id: str,
    watch_spreadsheet_id: str,
) -> dict[str, Any]:
    results: dict[str, Any] = {"work": {}, "watch": {}}
    calc_tabs = ["계산_계약명세", "계산_기준일투자자", "계산_기준일계약명세", "검증", "실행로그"]
    work_tabs = ["README", *raw_sheet_names(), *calc_tabs, *WORK_REPORT_SHEETS]
    watch_tabs = ["README", *raw_sheet_names(), *calc_tabs, *WATCH_REPORT_SHEETS]

    client.ensure_sheets(work_spreadsheet_id, work_tabs)
    client.ensure_sheets(watch_spreadsheet_id, watch_tabs)

    _write_common_tabs(client, work_spreadsheet_id, "업무보고서", "work")
    _write_common_tabs(client, watch_spreadsheet_id, "상시감시보고서", "watch")

    if WORKBOOK_PATH.exists():
        workbook_values = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True, read_only=True)
        workbook_styles = openpyxl.load_workbook(WORKBOOK_PATH, data_only=False, read_only=False)
        _write_templates(client, work_spreadsheet_id, workbook_values, workbook_styles, WORK_REPORT_SHEETS)
        _write_templates(client, watch_spreadsheet_id, workbook_values, workbook_styles, WATCH_REPORT_SHEETS)
        results["work"]["templateSheets"] = WORK_REPORT_SHEETS
        results["watch"]["templateSheets"] = WATCH_REPORT_SHEETS
    else:
        results["warning"] = f"템플릿 원본 xlsx를 찾지 못했습니다: {WORKBOOK_PATH}"

    return results


def update_report_templates_only(
    client: GoogleSheetsClient,
    work_spreadsheet_id: str,
    watch_spreadsheet_id: str,
) -> dict[str, Any]:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"템플릿 원본 xlsx를 찾지 못했습니다: {WORKBOOK_PATH}")
    workbook_values = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True, read_only=True)
    workbook_styles = openpyxl.load_workbook(WORKBOOK_PATH, data_only=False, read_only=False)
    _write_templates(client, work_spreadsheet_id, workbook_values, workbook_styles, WORK_REPORT_SHEETS)
    _write_templates(client, watch_spreadsheet_id, workbook_values, workbook_styles, WATCH_REPORT_SHEETS)
    return {"work": WORK_REPORT_SHEETS, "watch": WATCH_REPORT_SHEETS}


def raw_sheet_names() -> list[str]:
    return list(dict.fromkeys(spec.sheet_name for spec in SOURCE_SPECS.values()))


def _write_common_tabs(client: GoogleSheetsClient, spreadsheet_id: str, report_name: str, report_target: str) -> None:
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    readme = [
        [f"{report_name} 온라인 작성 문서"],
        [""],
        ["항목", "값"],
        ["문서구분", report_name],
        ["기준년", ""],
        ["기준월", ""],
        ["계산위치", "온라인 Apps Script"],
        ["로컬파이썬산출", "사용안함"],
        ["원천 데이터 탭", ", ".join(raw_sheet_names())],
        ["템플릿 생성일", now],
        [""],
        ["이 문서는 Codex 업로더가 관리하는 탭 구조입니다."],
        [""],
        ["업로드 흐름"],
        ["1", "웹 UI에서 서비스 계정 JSON과 업무/상시 문서 ID를 확인합니다."],
        ["2", "전체계약, 기준일투자자, 기준일계약 파일을 드래그 앤 드롭합니다."],
        ["3", "파일별 컬럼 체크박스에서 업로드할 컬럼을 확인합니다."],
        ["4", "업로드하면 선택한 컬럼만 원천 데이터 탭에 반영됩니다."],
        [""],
        ["주의", "보고서 탭은 엑셀 작성기의 출력 양식을 옮긴 템플릿입니다. 세부 계산 로직은 Google Apps Script에서 실행됩니다."],
    ]
    client.write_values(spreadsheet_id, "README", readme)

    for key in ("contract_all", "investor_snapshot", "contract_snapshot_current"):
        spec = SOURCE_SPECS[key]
        columns = list(default_columns_for_target(key, report_target))
        client.write_values(spreadsheet_id, spec.sheet_name, [columns])
        client.resize_and_freeze(spreadsheet_id, spec.sheet_name, rows=1, cols=len(columns))

def _write_templates(
    client: GoogleSheetsClient,
    spreadsheet_id: str,
    workbook_values: openpyxl.Workbook,
    workbook_styles: openpyxl.Workbook,
    sheet_names: list[str],
) -> None:
    for sheet_name in sheet_names:
        if sheet_name not in workbook_values.sheetnames:
            client.write_values(spreadsheet_id, sheet_name, [[sheet_name], ["원본 엑셀에서 시트를 찾지 못했습니다."]])
            continue
        ws_values = workbook_values[sheet_name]
        ws_styles = workbook_styles[sheet_name]
        rows = _worksheet_values(ws_values)
        client.write_values(spreadsheet_id, sheet_name, rows)
        _apply_worksheet_format(client, spreadsheet_id, sheet_name, ws_styles, len(rows), max((len(row) for row in rows), default=1))


def _worksheet_values(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[list[Any]]:
    max_row = ws.max_row
    max_col = ws.max_column
    values: list[list[Any]] = []
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
        out = [_normalize_cell(value) for value in row]
        values.append(out)

    while values and all(value == "" for value in values[-1]):
        values.pop()
    if not values:
        return [[ws.title]]
    last_col = 0
    for row in values:
        for idx, value in enumerate(row, start=1):
            if value != "":
                last_col = max(last_col, idx)
    return [row[:last_col] for row in values]


def _normalize_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        return _clean_excel_text(value)
    return value


def _clean_excel_text(value: str) -> str:
    cleaned = re.sub(r"_x000D_\s*", "", value)
    cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
    return cleaned.strip() if cleaned.startswith("_x000D_") else cleaned


def _apply_worksheet_format(
    client: GoogleSheetsClient,
    spreadsheet_id: str,
    sheet_name: str,
    ws: openpyxl.worksheet.worksheet.Worksheet,
    row_count: int,
    col_count: int,
) -> None:
    sheet_info = client.sheet_map(spreadsheet_id).get(sheet_name)
    if not sheet_info:
        return
    sheet_id = sheet_info.sheet_id
    requests: list[dict[str, Any]] = []

    requests.extend(_unmerge_requests(client, spreadsheet_id, sheet_id))
    requests.append(
        {
            "updateSheetProperties": {
                "properties": {
                    "sheetId": sheet_id,
                    "gridProperties": {
                        "rowCount": max(row_count, 100),
                        "columnCount": max(col_count, 20),
                        "frozenRowCount": 0,
                    },
                },
                "fields": "gridProperties(rowCount,columnCount,frozenRowCount)",
            }
        }
    )
    requests.append(
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id},
                "cell": {"userEnteredFormat": {}},
                "fields": "userEnteredFormat",
            }
        }
    )

    requests.extend(_dimension_requests(sheet_id, ws, row_count, col_count))
    requests.extend(_merge_requests(sheet_id, ws))
    requests.extend(_format_requests(sheet_id, ws, row_count, col_count))

    for idx in range(0, len(requests), 900):
        client._batch_update(spreadsheet_id, requests[idx : idx + 900])


def _unmerge_requests(client: GoogleSheetsClient, spreadsheet_id: str, sheet_id: int) -> list[dict[str, Any]]:
    meta = client.service.spreadsheets().get(
        spreadsheetId=spreadsheet_id,
        fields="sheets(properties(sheetId),merges)",
    ).execute()
    requests: list[dict[str, Any]] = []
    for sheet in meta.get("sheets", []):
        if sheet.get("properties", {}).get("sheetId") != sheet_id:
            continue
        for merge in sheet.get("merges", []):
            requests.append({"unmergeCells": {"range": merge}})
    return requests


def _dimension_requests(sheet_id: int, ws: openpyxl.worksheet.worksheet.Worksheet, row_count: int, col_count: int) -> list[dict[str, Any]]:
    requests: list[dict[str, Any]] = []
    for col_idx in range(1, col_count + 1):
        letter = openpyxl.utils.get_column_letter(col_idx)
        width = ws.column_dimensions[letter].width
        if width:
            pixel_size = max(32, min(320, int(width * 7 + 8)))
            requests.append(
                {
                    "updateDimensionProperties": {
                        "range": {"sheetId": sheet_id, "dimension": "COLUMNS", "startIndex": col_idx - 1, "endIndex": col_idx},
                        "properties": {"pixelSize": pixel_size},
                        "fields": "pixelSize",
                    }
                }
            )
    for row_idx in range(1, row_count + 1):
        height = ws.row_dimensions[row_idx].height
        if height:
            pixel_size = max(18, min(180, int(height * 1.35)))
            requests.append(
                {
                    "updateDimensionProperties": {
                        "range": {"sheetId": sheet_id, "dimension": "ROWS", "startIndex": row_idx - 1, "endIndex": row_idx},
                        "properties": {"pixelSize": pixel_size},
                        "fields": "pixelSize",
                    }
                }
            )
    return requests


def _merge_requests(sheet_id: int, ws: openpyxl.worksheet.worksheet.Worksheet) -> list[dict[str, Any]]:
    requests: list[dict[str, Any]] = []
    for merged in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merged))
        requests.append(
            {
                "mergeCells": {
                    "range": {
                        "sheetId": sheet_id,
                        "startRowIndex": min_row - 1,
                        "endRowIndex": max_row,
                        "startColumnIndex": min_col - 1,
                        "endColumnIndex": max_col,
                    },
                    "mergeType": "MERGE_ALL",
                }
            }
        )
    return requests


def _format_requests(sheet_id: int, ws: openpyxl.worksheet.worksheet.Worksheet, row_count: int, col_count: int) -> list[dict[str, Any]]:
    requests: list[dict[str, Any]] = []
    for row_idx in range(1, row_count + 1):
        current_key: str | None = None
        current_format: dict[str, Any] | None = None
        start_col = 1
        for col_idx in range(1, col_count + 2):
            fmt = _cell_format(ws.cell(row_idx, col_idx)) if col_idx <= col_count else None
            key = json.dumps(fmt, sort_keys=True, ensure_ascii=False) if fmt else None
            if key != current_key:
                if current_key and current_format:
                    requests.append(_repeat_cell_request(sheet_id, row_idx, start_col, col_idx, current_format))
                current_key = key
                current_format = fmt
                start_col = col_idx
    return requests


def _repeat_cell_request(sheet_id: int, row_idx: int, start_col: int, end_col: int, fmt: dict[str, Any]) -> dict[str, Any]:
    return {
        "repeatCell": {
            "range": {
                "sheetId": sheet_id,
                "startRowIndex": row_idx - 1,
                "endRowIndex": row_idx,
                "startColumnIndex": start_col - 1,
                "endColumnIndex": end_col - 1,
            },
            "cell": {"userEnteredFormat": fmt},
            "fields": "userEnteredFormat",
        }
    }


def _cell_format(cell: openpyxl.cell.cell.Cell) -> dict[str, Any] | None:
    fmt: dict[str, Any] = {}
    text_format: dict[str, Any] = {}
    if cell.font:
        if cell.font.bold:
            text_format["bold"] = True
        if cell.font.italic:
            text_format["italic"] = True
        if cell.font.sz:
            text_format["fontSize"] = int(cell.font.sz)
        font_color = _color(cell.font.color)
        if font_color:
            text_format["foregroundColor"] = font_color
    if text_format:
        fmt["textFormat"] = text_format

    bg = _fill_color(cell)
    if bg:
        fmt["backgroundColor"] = bg

    if cell.alignment:
        if cell.alignment.horizontal:
            fmt["horizontalAlignment"] = _horizontal(cell.alignment.horizontal)
        if cell.alignment.vertical:
            fmt["verticalAlignment"] = _vertical(cell.alignment.vertical)
        if cell.alignment.wrap_text:
            fmt["wrapStrategy"] = "WRAP"

    if cell.number_format and cell.number_format != "General":
        fmt["numberFormat"] = {"type": _number_type(cell.number_format), "pattern": cell.number_format}

    borders = _borders(cell)
    if borders:
        fmt["borders"] = borders

    return fmt or None


def _fill_color(cell: openpyxl.cell.cell.Cell) -> dict[str, float] | None:
    fill = cell.fill
    if not fill or fill.fill_type not in {"solid", "gray125"}:
        return None
    return _color(fill.fgColor)


def _color(color: Any) -> dict[str, float] | None:
    if color is None:
        return None
    rgb = None
    if getattr(color, "type", None) == "rgb" and color.rgb:
        rgb = color.rgb[-6:]
    elif getattr(color, "type", None) == "indexed" and color.indexed is not None:
        try:
            indexed = openpyxl.styles.colors.COLOR_INDEX[color.indexed]
            rgb = indexed[-6:] if indexed else None
        except Exception:
            rgb = None
    if not rgb or rgb in {"000000", "FFFFFF"} and getattr(color, "type", None) != "rgb":
        return None
    try:
        return {
            "red": int(rgb[0:2], 16) / 255,
            "green": int(rgb[2:4], 16) / 255,
            "blue": int(rgb[4:6], 16) / 255,
        }
    except Exception:
        return None


def _horizontal(value: str) -> str:
    return {
        "center": "CENTER",
        "left": "LEFT",
        "right": "RIGHT",
        "justify": "LEFT",
    }.get(value, "CENTER")


def _vertical(value: str) -> str:
    return {
        "center": "MIDDLE",
        "top": "TOP",
        "bottom": "BOTTOM",
    }.get(value, "MIDDLE")


def _number_type(pattern: str) -> str:
    lower = pattern.lower()
    if "%" in lower:
        return "PERCENT"
    if any(token in lower for token in ["yy", "mm", "dd"]):
        return "DATE"
    if any(token in lower for token in ["#", "0"]):
        return "NUMBER"
    return "TEXT"


def _borders(cell: openpyxl.cell.cell.Cell) -> dict[str, Any] | None:
    result: dict[str, Any] = {}
    for excel_side, sheet_side in [("left", "left"), ("right", "right"), ("top", "top"), ("bottom", "bottom")]:
        side = getattr(cell.border, excel_side)
        if side and side.style:
            result[sheet_side] = {
                "style": _border_style(side.style),
                "color": _color(side.color) or {"red": 0.75, "green": 0.75, "blue": 0.75},
            }
    return result or None


def _border_style(style: str) -> str:
    if style in {"medium", "thick", "double"}:
        return "SOLID_MEDIUM"
    if style in {"dashed", "dashDot", "dashDotDot"}:
        return "DASHED"
    if style in {"dotted", "hair"}:
        return "DOTTED"
    return "SOLID"
