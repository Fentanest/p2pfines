from __future__ import annotations

from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any
import re
import shutil
import zipfile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from .config import OUTPUT_DIR
from .paths import RESOURCE_ROOT
from .local_watch_calculator import apply_watch_calculations
from .local_work_calculator import apply_work_calculations


TEMPLATE_CANDIDATES = [
    RESOURCE_ROOT / "june.xlsx",
    RESOURCE_ROOT / "업무보고서 작성기_v5.4_통합.xlsx",
]

RAW_CONTRACT_ALL = "계약명세"
RAW_INVESTOR_TEMPLATE = "기준일투자자데이터"
RAW_CONTRACT_SNAPSHOT_TEMPLATE = "기준일명세서"

LOCAL_SHEET_BY_SOURCE_KEY = {
    "contract_all": RAW_CONTRACT_ALL,
    "investor_snapshot": RAW_INVESTOR_TEMPLATE,
    "contract_snapshot": RAW_CONTRACT_SNAPSHOT_TEMPLATE,
    "contract_snapshot_prev": RAW_CONTRACT_SNAPSHOT_TEMPLATE,
    "contract_snapshot_current": RAW_CONTRACT_SNAPSHOT_TEMPLATE,
}

WORK_OUTPUT_SHEETS = [
    RAW_CONTRACT_ALL,
    RAW_CONTRACT_SNAPSHOT_TEMPLATE,
    RAW_INVESTOR_TEMPLATE,
    "PL302",
    "PL303",
    "PL305",
    "PL307",
    "PL309",
    "PL313",
    "PL315",
    "PL317",
    "PL319",
    "PL320",
]

WATCH_OUTPUT_SHEETS = [
    RAW_CONTRACT_ALL,
    RAW_CONTRACT_SNAPSHOT_TEMPLATE,
    RAW_INVESTOR_TEMPLATE,
    "(상시)1-1",
    "(상시)1-3",
    "(상시)2-2",
    "(상시)2-3",
    "(상시)3-1",
    "(상시)3-2",
]


@dataclass
class LocalExcelResult:
    path: Path
    file_name: str
    warnings: list[str] = field(default_factory=list)


class LocalExcelError(RuntimeError):
    pass


def generate_local_excel(
    prepared: list[dict[str, Any]],
    target: str,
    year: int,
    month: int,
) -> LocalExcelResult:
    template = _find_template()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    target_label = "업무보고서" if target == "work" else "상시감시보고서"
    value_only_suffix = "_값전용" if target in {"work", "watch"} else ""
    file_name = f"{target_label}{value_only_suffix}_{year}-{month:02d}_{timestamp}.xlsx"
    output_path = OUTPUT_DIR / file_name

    shutil.copy2(template, output_path)
    wb = load_workbook(output_path)
    warnings: list[str] = []

    _validate_unique_periods(prepared)
    grouped = _group_prepared(prepared)
    copy_template_formulas = target not in {"work", "watch"}
    for sheet_name, df in grouped.items():
        if sheet_name not in wb.sheetnames:
            warnings.append(f"템플릿에 원천 시트가 없어 건너뜀: {sheet_name}")
            continue
        _write_source_dataframe(wb[sheet_name], df, copy_formulas=copy_template_formulas)

    _write_period_cells(wb, target, year, month)
    locally_calculated = False
    if target == "work":
        apply_work_calculations(wb, grouped, year, month)
        locally_calculated = True
    elif target == "watch":
        apply_watch_calculations(wb, grouped, year, month)
        locally_calculated = True
    _prune_output_sheets(wb, target)
    if locally_calculated:
        _strip_formula_cells(wb)
    else:
        _set_recalculate_on_open(wb)
    wb.save(output_path)
    if locally_calculated:
        _remove_formula_xml_tags(output_path)

    if not locally_calculated:
        try:
            _recalculate_with_excel(output_path, target)
        except Exception as exc:
            warnings.append(f"Excel 백그라운드 재계산 실패: {exc}. 파일은 수식 포함 상태로 저장했습니다.")

    return LocalExcelResult(path=output_path, file_name=file_name, warnings=warnings)


def _find_template() -> Path:
    for path in TEMPLATE_CANDIDATES:
        if path.exists():
            return path
    raise LocalExcelError("로컬 엑셀 생성용 템플릿을 찾지 못했습니다. june.xlsx 또는 업무보고서 작성기 파일을 프로젝트 루트에 두세요.")


def _group_prepared(prepared: list[dict[str, Any]]) -> dict[str, pd.DataFrame]:
    grouped: dict[str, list[pd.DataFrame]] = {}
    for item in prepared:
        sheet_name = LOCAL_SHEET_BY_SOURCE_KEY.get(item["sourceKey"])
        if not sheet_name:
            continue
        grouped.setdefault(sheet_name, []).append(item["df"])

    out: dict[str, pd.DataFrame] = {}
    for sheet_name, frames in grouped.items():
        combined = pd.concat(frames, ignore_index=True, sort=False).fillna("")
        out[sheet_name] = _sort_dataframe_by_period(combined)
    return out


def _validate_unique_periods(prepared: list[dict[str, Any]]) -> None:
    seen: dict[tuple[str, tuple[int, int]], str] = {}
    for item in prepared:
        sheet_name = LOCAL_SHEET_BY_SOURCE_KEY.get(item["sourceKey"])
        period_key = item.get("periodKey")
        if not sheet_name or not period_key:
            continue
        key = (sheet_name, period_key)
        original_name = str(item.get("originalName") or "")
        if key in seen:
            year, month = period_key
            raise LocalExcelError(
                f"{sheet_name} {year}년 {month}월 파일이 두 번 포함되었습니다: "
                f"{seen[key]}, {original_name}. 중복 월을 제거한 뒤 다시 생성해주세요."
            )
        seen[key] = original_name


def _sort_dataframe_by_period(df: pd.DataFrame) -> pd.DataFrame:
    if "기준년" not in df.columns or "기준월" not in df.columns:
        return df
    sorted_df = df.copy()
    sorted_df["_period_year_sort"] = pd.to_numeric(sorted_df["기준년"], errors="coerce").fillna(9999)
    sorted_df["_period_month_sort"] = pd.to_numeric(sorted_df["기준월"], errors="coerce").fillna(99)
    sorted_df["_period_original_order"] = range(len(sorted_df))
    sorted_df = sorted_df.sort_values(
        by=["_period_year_sort", "_period_month_sort", "_period_original_order"],
        kind="mergesort",
    )
    return sorted_df.drop(columns=["_period_year_sort", "_period_month_sort", "_period_original_order"])


def _write_source_dataframe(ws, df: pd.DataFrame, copy_formulas: bool = True) -> None:
    if not copy_formulas:
        _write_compact_source_dataframe(ws, df)
        return

    header_row = 2
    data_start_row = header_row + 1
    max_col = ws.max_column
    row_template = _capture_row_template(ws, data_start_row, max_col)
    headers = {
        str(ws.cell(header_row, col).value).strip(): col
        for col in range(1, max_col + 1)
        if ws.cell(header_row, col).value not in (None, "")
    }

    _clear_data_area(ws, data_start_row)

    for row_offset, (_, record) in enumerate(df.iterrows()):
        row_idx = data_start_row + row_offset
        _apply_row_template(ws, row_idx, row_template, data_start_row, copy_formulas=copy_formulas)
        for column_name, value in record.items():
            col_idx = headers.get(str(column_name).strip())
            if not col_idx:
                continue
            ws.cell(row_idx, col_idx).value = _excel_value(value)


def _write_compact_source_dataframe(ws, df: pd.DataFrame) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))

    if ws.max_row:
        ws.delete_rows(1, ws.max_row)
    if ws.max_column:
        ws.delete_cols(1, ws.max_column)

    columns = [str(column) for column in df.columns]
    if not columns:
        ws.cell(1, 1).value = ""
        return

    for col_idx, column_name in enumerate(columns, start=1):
        cell = ws.cell(1, col_idx)
        cell.value = column_name
        cell.number_format = "@"
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(len(column_name) + 2, 10), 24)

    for row_offset, (_, record) in enumerate(df.iterrows(), start=2):
        for col_idx, column_name in enumerate(columns, start=1):
            ws.cell(row_offset, col_idx).value = _excel_value(record.get(column_name, ""))

    last_cell = f"{get_column_letter(len(columns))}{max(len(df) + 1, 1)}"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last_cell}"


def _capture_row_template(ws, row_idx: int, max_col: int) -> dict[int, dict[str, Any]]:
    out: dict[int, dict[str, Any]] = {}
    for col_idx in range(1, max_col + 1):
        cell = ws.cell(row_idx, col_idx)
        out[col_idx] = {
            "value": cell.value,
            "style": copy(cell._style),
            "number_format": cell.number_format,
            "fill": copy(cell.fill),
            "font": copy(cell.font),
            "alignment": copy(cell.alignment),
            "border": copy(cell.border),
            "protection": copy(cell.protection),
        }
    return out


def _clear_data_area(ws, start_row: int) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            cell.value = None


def _apply_row_template(
    ws,
    row_idx: int,
    template: dict[int, dict[str, Any]],
    template_row: int,
    copy_formulas: bool = True,
) -> None:
    for col_idx, spec in template.items():
        cell = ws.cell(row_idx, col_idx)
        if isinstance(cell, MergedCell):
            continue
        cell._style = copy(spec["style"])
        cell.number_format = spec["number_format"]
        cell.fill = copy(spec["fill"])
        cell.font = copy(spec["font"])
        cell.alignment = copy(spec["alignment"])
        cell.border = copy(spec["border"])
        cell.protection = copy(spec["protection"])

        value = spec["value"]
        if copy_formulas and isinstance(value, str) and value.startswith("="):
            origin = f"{get_column_letter(col_idx)}{template_row}"
            cell.value = Translator(value, origin=origin).translate_formula(cell.coordinate)


def _excel_value(value: Any) -> Any:
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _strip_formula_cells(wb) -> int:
    stripped = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    cell.value = None
                    stripped += 1
    return stripped


def _remove_formula_xml_tags(path: Path) -> int:
    tmp_path = path.with_suffix(path.suffix + ".tmp")
    tmp_path.unlink(missing_ok=True)
    removed = 0
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            name = item.filename
            if name == "xl/calcChain.xml":
                continue
            data = source.read(item)
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                text = data.decode("utf-8")
                text, count = re.subn(r"<f(?:\s[^>]*)?>.*?</f>", "", text, flags=re.DOTALL)
                removed += count
                text, count = re.subn(r"<f\s[^>]*/>", "", text, flags=re.DOTALL)
                removed += count
                data = text.encode("utf-8")
            elif name == "[Content_Types].xml":
                text = data.decode("utf-8")
                text = re.sub(r'<Override PartName="/xl/calcChain.xml"[^>]*/>', "", text)
                data = text.encode("utf-8")
            elif name == "xl/_rels/workbook.xml.rels":
                text = data.decode("utf-8")
                text = re.sub(r"<Relationship[^>]+calcChain[^>]*/>", "", text)
                data = text.encode("utf-8")
            target.writestr(item, data)
    tmp_path.replace(path)
    return removed


def _prune_output_sheets(wb, target: str) -> None:
    keep_order = WORK_OUTPUT_SHEETS if target == "work" else WATCH_OUTPUT_SHEETS
    keep = set(keep_order)
    for sheet_name in list(wb.sheetnames):
        if sheet_name not in keep and len(wb.sheetnames) > 1:
            del wb[sheet_name]
    for index, sheet_name in enumerate(keep_order):
        if sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            wb._sheets.remove(sheet)
            wb._sheets.insert(index, sheet)


def _write_period_cells(wb, target: str, year: int, month: int) -> None:
    if target == "work":
        sheet_names = ["PL302", "PL303", "PL305", "PL307", "PL309", "PL313", "PL315", "PL317", "PL319", "PL320"]
    else:
        sheet_names = ["(상시)1-1", "(상시)2-2", "(상시)3-1", "(상시)3-2"]
    for sheet_name in sheet_names:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            ws["A1"] = year
            ws["B1"] = month


def _set_recalculate_on_open(wb) -> None:
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass


def _recalculate_with_excel(path: Path, target: str) -> None:
    import win32com.client

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        excel.Calculation = -4135  # xlCalculationManual
    except Exception:
        pass
    workbook = None
    try:
        workbook = excel.Workbooks.Open(str(path.resolve()))
        for sheet_name in _calculation_sheet_order(target):
            try:
                workbook.Worksheets(sheet_name).Calculate()
            except Exception:
                pass
        excel.CalculateUntilAsyncQueriesDone()
        workbook.Save()
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        excel.Quit()


def _calculation_sheet_order(target: str) -> list[str]:
    common = [
        RAW_CONTRACT_ALL,
        RAW_CONTRACT_SNAPSHOT_TEMPLATE,
        RAW_INVESTOR_TEMPLATE,
    ]
    if target == "work":
        return [
            *common,
            "PL302",
            "PL303",
            "PL305",
            "PL307",
            "PL309",
            "PL313",
            "PL315",
            "PL317",
            "PL319",
            "PL320",
        ]
    return [
        *common,
        "월별실적",
        "(상시)1-1",
        "(상시)1-3",
        "(상시)3-1",
        "(상시)2-2",
        "(상시)2-3",
        "(상시)3-2",
    ]
