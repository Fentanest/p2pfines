from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta
from math import isfinite
from typing import Any, Iterable

import pandas as pd
from openpyxl.cell.cell import MergedCell


RAW_CONTRACT_ALL = "계약명세"
RAW_INVESTOR = "기준일투자자데이터"
RAW_SNAPSHOT = "기준일명세서"


def apply_work_calculations(wb, grouped: dict[str, pd.DataFrame], year: int, month: int) -> None:
    contract_all = _records(grouped.get(RAW_CONTRACT_ALL))
    investor = _records(grouped.get(RAW_INVESTOR))
    snapshots = _records(grouped.get(RAW_SNAPSHOT))

    contract_all_norm = _normalize_contract_all(contract_all)
    contract_by_no = {_text(row.get("계약No")): row for row in contract_all_norm if _text(row.get("계약No"))}
    investor_norm = _normalize_investor(investor, contract_by_no, year, month)
    snapshot_norm = _normalize_snapshots(snapshots, contract_by_no, year, month)

    _write_period(wb, year, month)
    _clear_work_outputs(wb)

    _write_product_values(wb, "PL302", 14, 5, _calc_pl302(investor_norm, year, month), 4)
    _write_product_values(wb, "PL303", 14, 5, _calc_pl303(investor_norm, year, month), 4)
    _calc_pl305(wb, investor_norm, contract_all_norm, year, month)
    _write_product_values(wb, "PL307", 14, 5, _calc_pl307(snapshot_norm, year, month), 4)
    _write_product_values(wb, "PL309", 14, 5, _calc_pl309(snapshot_norm, year, month), 4)
    _write_real_estate_values(wb, "PL313", _calc_pl313(snapshot_norm, year, month), 3)
    _write_product_values(wb, "PL315", 14, 5, _calc_pl315(snapshot_norm, year, month), 4)
    _write_pl317(wb, _calc_pl317(snapshot_norm, year, month))
    _write_pl319(wb, _calc_pl319(snapshot_norm, year, month))


def _records(df: pd.DataFrame | None) -> list[dict[str, Any]]:
    if df is None or df.empty:
        return []
    clean = df.fillna("")
    return clean.to_dict("records")


def _normalize_contract_all(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        item = dict(row)
        _apply_adjusted_status(item)
        item["상품명_정리"] = _clean_product_name(item.get("상품명"))
        out.append(item)
    return out


def _normalize_investor(
    rows: list[dict[str, Any]],
    contract_by_no: dict[str, dict[str, Any]],
    year: int,
    month: int,
) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        item = dict(row)
        item["기준년"] = _value_or(item.get("기준년"), year)
        item["기준월"] = _value_or(item.get("기준월"), month)
        contract = contract_by_no.get(_text(item.get("계약번호")), {})
        item["대출자구분"] = _value_or(item.get("대출자구분"), contract.get("대출자구분"))
        item["투자자구분_정리"] = _normalize_investor_type(item.get("투자자구분"))
        item["상품분류_정리"] = _normalize_product(item.get("세부상품") or contract.get("세부상품"))
        item["투자금액_숫자"] = _number(item.get("투자금액"))
        item["투자잔액_숫자"] = _number(item.get("투자잔액"))
        out.append(item)
    return out


def _normalize_snapshots(
    rows: list[dict[str, Any]],
    contract_by_no: dict[str, dict[str, Any]],
    year: int,
    month: int,
) -> list[dict[str, Any]]:
    prepared = []
    for row in rows:
        item = dict(row)
        contract = contract_by_no.get(_text(item.get("계약No")), {})
        _apply_adjusted_status(item, contract)
        item["기준년"] = _value_or(item.get("기준년"), year)
        item["기준월"] = _value_or(item.get("기준월"), month)
        item["잔액_숫자"] = _number(item.get("잔액"))
        item["총대출액_숫자"] = _number(item.get("총대출액"))
        item["LTV_숫자"] = _number(item.get("LTV"))
        item["기간_숫자"] = _number(item.get("기간"))
        item["월-계약No"] = "-".join(_text(item.get(key)) for key in ("기준년", "기준월", "계약No"))
        item["월-계약No-상태"] = "-".join(_text(item.get(key)) for key in ("기준년", "기준월", "계약No", "상태"))
        item["월-주민등록번호"] = "-".join(_text(item.get(key)) for key in ("기준년", "기준월", "주민등록"))
        item["상품분류_정리"] = _normalize_product(item.get("세부상품"))
        item["PL309_계산"] = _value_or(item.get("PL309"), _amount_band(item["잔액_숫자"]))
        item["PL313_계산"] = _value_or(item.get("PL313"), _ltv_band(item["LTV_숫자"]))
        item["PL315_계산"] = _value_or(item.get("PL315"), _term_band(item["기간_숫자"]))
        item["PL317_계산"] = _value_or(item.get("PL317"), _overdue_band(item))
        prepared.append(item)

    by_person_month = _group_by(
        prepared,
        lambda row: "|".join(_text(row.get(key)) for key in ("기준년", "기준월", "주민등록")),
    )
    by_contract_month = _group_by(
        prepared,
        lambda row: "|".join(_text(row.get(key)) for key in ("기준년", "기준월", "계약No")),
    )

    for row in prepared:
        ym = _period(_number(row.get("기준년")), _number(row.get("기준월")))
        prev = _add_period(ym, -1)
        next_period = _add_period(ym, 1)
        person = _text(row.get("주민등록"))
        current_rows = by_person_month.get("|".join(map(_text, [ym["year"], ym["month"], person])), [])
        prev_rows = by_person_month.get("|".join(map(_text, [prev["year"], prev["month"], person])), [])
        next_rows = by_person_month.get("|".join(map(_text, [next_period["year"], next_period["month"], person])), [])

        max_balance = _max(row.get("잔액_숫자") for row in current_rows)
        max_ltv = _max(row.get("LTV_숫자") for row in current_rows)
        max_term = _max(row.get("기간_숫자") for row in current_rows)
        max_contract_no = _max(row.get("계약No") for row in current_rows)
        max_balance_contract_no = _max_contract_for_value(current_rows, "잔액_숫자", max_balance)
        max_ltv_contract_no = _max_contract_for_value(current_rows, "LTV_숫자", max_ltv)
        max_term_contract_no = _max_contract_for_value(current_rows, "기간_숫자", max_term)

        row["pl307 신규차입자"] = 1 if _sum(prev_rows, "잔액_숫자") == 0 and _number(row.get("계약No")) == max_contract_no else 0
        prev_contract_key = "|".join(map(_text, [prev["year"], prev["month"], row.get("계약No")]))
        prev_contract_rows = by_contract_month.get(prev_contract_key, [])
        row["pl307 신규취급액"] = 0 if prev_contract_rows else max(0, _number(row.get("총대출액_숫자")))

        row["pl309 중복확인"] = 1 if row["잔액_숫자"] != 0 and row["잔액_숫자"] == max_balance and _number(row.get("계약No")) == max_balance_contract_no else 0
        row["pl313 중복확인"] = 1 if row["잔액_숫자"] != 0 and row["LTV_숫자"] == max_ltv and _number(row.get("계약No")) == max_ltv_contract_no else 0
        row["pl315 중복확인"] = 1 if row["잔액_숫자"] != 0 and row["기간_숫자"] == max_term and _number(row.get("계약No")) == max_term_contract_no else 0

        repayment_period = prev if row["잔액_숫자"] == 0 else ym
        row["pl307 상환기준년"] = repayment_period["year"]
        row["pl307 상환기준월"] = repayment_period["month"]
        row["pl307 완제차입자"] = 1 if (
            (next_rows and row["잔액_숫자"] == 0)
            or (not next_rows and _number(row.get("계약No")) == max_contract_no)
        ) else 0
        next_contract_key = "|".join(map(_text, [next_period["year"], next_period["month"], row.get("계약No")]))
        next_contract_rows = by_contract_month.get(next_contract_key, [])
        next_balance = _number(next_contract_rows[0].get("잔액_숫자")) if next_contract_rows else 0
        row["pl307 상환액"] = row["잔액_숫자"] - next_balance + (row["총대출액_숫자"] if row["잔액_숫자"] == 0 else 0)

    return prepared


def _apply_adjusted_status(row: dict[str, Any], contract: dict[str, Any] | None = None) -> None:
    raw_status = _text(row.get("상태"))
    contract = contract or {}
    overdue_days = _overdue_days(row, contract)
    adjusted = _adjusted_status(raw_status, overdue_days)
    row["상태_원천"] = raw_status
    row["연체일_숫자"] = overdue_days if overdue_days is not None else ""
    row["상태분류_계산"] = adjusted
    row["상태"] = adjusted


def _adjusted_status(raw_status: str, overdue_days: float | None) -> str:
    if "매각" in raw_status or "완제" in raw_status or "완납" in raw_status:
        return raw_status
    if overdue_days is not None:
        return "연체" if overdue_days >= 30 else "정상"
    if "연체" in raw_status:
        return "연체"
    return raw_status or "정상"


def _overdue_days(row: dict[str, Any], contract: dict[str, Any]) -> float | None:
    for key in ("연체일", "최장연체일수"):
        value = row.get(key)
        if value not in (None, ""):
            return _number(value)
    for key in ("연체일", "최장연체일수"):
        value = contract.get(key)
        if value not in (None, ""):
            return _number(value)
    return None


def _calc_pl302(investor: list[dict[str, Any]], year: int, month: int) -> list[list[Any]]:
    rows = []
    for definition in _product_rows():
        filtered = [
            row for row in investor
            if _same_period(row, year, month) and _match_borrower(row, definition["borrower"]) and _match_product(row, definition["products"])
        ]
        result: list[Any] = [definition["code"], definition["borrowerLabel"], definition["productLabel"], definition.get("detailLabel", "")]
        for investor_type in _investor_types():
            typed = [row for row in filtered if row.get("투자자구분_정리") == investor_type["key"]]
            result.extend([_unique_count_nonzero(typed, "주민등록번호"), _sum(typed, "투자잔액_숫자")])
        result.extend([_sum_every_other(result, 4), _sum_every_other(result, 5)])
        rows.append(result)
    return rows


def _calc_pl303(investor: list[dict[str, Any]], year: int, month: int) -> list[list[Any]]:
    bands = _invest_amount_bands()
    current_rows = [row for row in investor if _same_period(row, year, month)]
    period_sums = _group_sum(current_rows, "주민등록번호", "투자잔액_숫자")
    rows = []
    for definition in _product_rows():
        filtered = [
            row for row in investor
            if _same_period(row, year, month) and _match_borrower(row, definition["borrower"]) and _match_product(row, definition["products"])
        ]
        ids = _unique_values(row.get("주민등록번호") for row in filtered)
        amounts_for_ids = [period_sums.get(investor_id, 0) for investor_id in ids]
        result: list[Any] = [definition["code"], definition["borrowerLabel"], definition["productLabel"], definition.get("detailLabel", "")]
        for band in bands:
            amounts = [value for value in amounts_for_ids if value > band["min"] and value <= band["max"]]
            result.extend([len(amounts), sum(amounts)])
        result.extend([len(amounts_for_ids), sum(amounts_for_ids)])
        rows.append(result)
    return rows


def _calc_pl305(wb, investor: list[dict[str, Any]], contract_all: list[dict[str, Any]], year: int, month: int) -> None:
    contract_by_product = {}
    for row in contract_all:
        key = _clean_product_name(row.get("상품명"))
        if key:
            contract_by_product[key] = row
    filtered = [
        row for row in investor
        if _same_period(row, year, month) and row.get("투자자구분_정리") in {"법인", "전문"}
    ]
    rows = []
    for row in filtered:
        product = _clean_product_name(row.get("상품명"))
        contract = contract_by_product.get(product, {})
        investor_type = "법인투자자" if row.get("투자자구분_정리") == "법인" else "개인전문투자자"
        rows.append([
            product,
            _number(contract.get("한도") or contract.get("총대출") or 0),
            row.get("세부상품") or contract.get("세부상품") or "",
            _yyyymmdd(contract.get("계약일")),
            _yyyymmdd(contract.get("만기일")),
            _mask_name(row.get("이름")) if investor_type == "개인전문투자자" else row.get("이름"),
            investor_type,
            _business_no(row.get("주민등록번호")) if investor_type == "법인투자자" else "",
            _number(row.get("투자금액_숫자")),
        ])
    _write_rows(wb, "PL305", 8, 2, rows, 9)
    _write_pl305_investor_filter(wb, filtered)


def _write_pl305_investor_filter(wb, rows: list[dict[str, Any]]) -> None:
    if "PL305" not in wb.sheetnames:
        return
    ws = wb["PL305"]
    _clear_range(ws, 6, 13, max(ws.max_row - 5, 1), 12)
    headers = ["주민등록번호", "상품명", "세부상품", "이름", "투자자구분", "투자금액"]
    _write_matrix(ws, 6, 13, [["투자자리스트 필터링", "", "", "", "", ""]])
    ws.cell(6, 18).value = _sum(rows, "투자금액_숫자")
    _write_matrix(ws, 7, 13, [headers])
    if rows:
        _write_matrix(ws, 8, 13, [[
            row.get("주민등록번호") or "",
            _clean_product_name(row.get("상품명")),
            row.get("세부상품") or "",
            row.get("이름") or "",
            row.get("투자자구분") or "",
            _number(row.get("투자금액_숫자")),
        ] for row in rows])


def _calc_pl307(snapshots: list[dict[str, Any]], year: int, month: int) -> list[list[Any]]:
    ym = _period(year, month)
    prev = _add_period(ym, -1)
    has_prev_period = any(_same_period(row, prev["year"], prev["month"]) for row in snapshots)
    rows = []
    for definition in _product_rows():
        prev_rows = [
            row for row in snapshots
            if has_prev_period and _same_period(row, prev["year"], prev["month"]) and _match_borrower(row, definition["borrower"]) and _match_product(row, definition["products"])
        ]
        current_rows = [
            row for row in snapshots
            if _same_period(row, year, month) and _match_borrower(row, definition["borrower"]) and _match_product(row, definition["products"])
        ]
        repay_rows = [
            row for row in snapshots
            if has_prev_period and _number(row.get("pl307 상환기준년")) == prev["year"] and _number(row.get("pl307 상환기준월")) == prev["month"]
            and _match_borrower(row, definition["borrower"]) and _match_product(row, definition["products"])
        ]
        new_borrower_rows = [row for row in current_rows if _number(row.get("pl307 신규차입자")) > 0] if has_prev_period else []
        completed_rows = [row for row in repay_rows if _number(row.get("pl307 완제차입자")) > 0] if has_prev_period else []
        prev_count = len([row for row in prev_rows if _number(row.get("pl315 중복확인")) > 0])
        prev_amount = _sum(prev_rows, "잔액_숫자")
        new_count = len(new_borrower_rows)
        new_amount = _sum(current_rows, "pl307 신규취급액")
        repay_count = len(completed_rows)
        repay_amount = _sum(repay_rows, "pl307 상환액")
        current_count = prev_count + new_count - repay_count if has_prev_period else _unique_count([row for row in current_rows if _number(row.get("잔액_숫자")) != 0], "주민등록")
        current_amount = prev_amount + new_amount - repay_amount if has_prev_period else _sum(current_rows, "잔액_숫자")
        rows.append([
            definition["code"], definition["borrowerLabel"], definition["productLabel"], definition.get("detailLabel", ""),
            prev_count, prev_amount, new_count, new_amount, repay_count, repay_amount, current_count, current_amount,
        ])
    return rows


def _calc_pl309(snapshots: list[dict[str, Any]], year: int, month: int) -> list[list[Any]]:
    rows = []
    for definition in _product_rows():
        filtered = [
            row for row in snapshots
            if _same_period(row, year, month) and _match_borrower(row, definition["borrower"]) and _match_product(row, definition["products"])
        ]
        result: list[Any] = [definition["code"], definition["borrowerLabel"], definition["productLabel"], definition.get("detailLabel", "")]
        for band in _loan_amount_bands():
            band_rows = [row for row in filtered if (row.get("PL309_계산") or _amount_band(row.get("잔액_숫자"))) == band["label"]]
            result.extend([len([row for row in band_rows if _number(row.get("pl309 중복확인")) > 0]), _sum(band_rows, "잔액_숫자")])
        result.extend([_sum_every_other(result, 4), _sum_every_other(result, 5)])
        rows.append(result)
    return rows


def _calc_pl313(snapshots: list[dict[str, Any]], year: int, month: int) -> list[list[Any]]:
    rows = []
    for definition in _real_estate_rows():
        filtered = [
            row for row in snapshots
            if _same_period(row, year, month) and _match_borrower(row, definition["borrower"]) and _match_product(row, definition["products"])
        ]
        result: list[Any] = [definition["code"], definition["borrowerLabel"], definition["productLabel"]]
        for band in _ltv_bands():
            band_rows = [row for row in filtered if (row.get("PL313_계산") or _ltv_band(row.get("LTV_숫자"))) == band["label"]]
            result.extend([len([row for row in band_rows if _number(row.get("pl313 중복확인")) > 0]), _sum(band_rows, "잔액_숫자")])
        result.extend([_sum_every_other(result, 3), _sum_every_other(result, 4)])
        rows.append(result)
    return rows


def _calc_pl315(snapshots: list[dict[str, Any]], year: int, month: int) -> list[list[Any]]:
    rows = []
    for definition in _product_rows():
        filtered = [
            row for row in snapshots
            if _same_period(row, year, month) and _match_borrower(row, definition["borrower"]) and _match_product(row, definition["products"])
        ]
        result: list[Any] = [definition["code"], definition["borrowerLabel"], definition["productLabel"], definition.get("detailLabel", "")]
        for band in _term_bands():
            band_rows = [row for row in filtered if (row.get("PL315_계산") or _term_band(row.get("기간_숫자"))) == band["label"]]
            result.extend([len([row for row in band_rows if _number(row.get("pl315 중복확인")) > 0]), _sum(band_rows, "잔액_숫자")])
        result.extend([_sum_every_other(result, 4), _sum_every_other(result, 5)])
        rows.append(result)
    return rows


def _calc_pl317(snapshots: list[dict[str, Any]], year: int, month: int) -> list[list[Any]]:
    rows = []
    for definition in _product_rows():
        filtered = [
            row for row in snapshots
            if _same_period(row, year, month) and _match_borrower(row, definition["borrower"]) and _match_product(row, definition["products"])
        ]
        normal_rows = [row for row in filtered if not _text(row.get("PL317_계산") or _overdue_band(row))]
        overdue_rows = [row for row in filtered if _text(row.get("PL317_계산") or _overdue_band(row))]
        result: list[Any] = [definition["code"], definition["borrowerLabel"], definition["productLabel"], definition.get("detailLabel", "")]
        result.extend([len([row for row in normal_rows if _number(row.get("pl315 중복확인")) > 0]), _sum(normal_rows, "잔액_숫자")])
        for band in _term_bands():
            band_rows = [row for row in overdue_rows if (row.get("PL317_계산") or _overdue_band(row)) == band["label"]]
            result.extend([len([row for row in band_rows if _number(row.get("pl315 중복확인")) > 0]), _sum(band_rows, "잔액_숫자")])
        overdue_count = _number(result[6]) + _number(result[8]) + _number(result[10]) + _number(result[12]) + _number(result[14])
        overdue_amount = _number(result[7]) + _number(result[9]) + _number(result[11]) + _number(result[13]) + _number(result[15])
        total_count = _number(result[4]) + overdue_count
        total_amount = _number(result[5]) + overdue_amount
        result.extend([overdue_count, overdue_amount, overdue_count, overdue_amount, total_count, total_amount, overdue_amount / total_amount * 100 if total_amount else 0])
        rows.append(result)
    return rows


def _calc_pl319(snapshots: list[dict[str, Any]], year: int, month: int) -> list[list[Any]]:
    rows = []
    for definition in _product_rows():
        filtered = [
            row for row in snapshots
            if _same_period(row, year, month) and _match_borrower(row, definition["borrower"]) and _match_product(row, definition["products"])
        ]
        balance = _sum(filtered, "잔액_숫자")
        interest = sum(_number(row.get("잔액_숫자")) * _number(row.get("이율")) / 100 for row in filtered) / 12
        rates = [_number(row.get("이율")) for row in filtered if _number(row.get("이율")) > 0]
        rows.append([
            definition["code"], definition["borrowerLabel"], definition["productLabel"], definition.get("detailLabel", ""),
            balance, interest, interest * 12 / balance * 100 if balance else 0, max(rates) if rates else 0, min(rates) if rates else 0,
        ])
    return rows


def _write_period(wb, year: int, month: int) -> None:
    for sheet_name in ["PL302", "PL303", "PL305", "PL307", "PL309", "PL313", "PL315", "PL317", "PL319", "PL320"]:
        if sheet_name in wb.sheetnames:
            wb[sheet_name]["A1"] = year
            wb[sheet_name]["B1"] = month


def _clear_work_outputs(wb) -> None:
    specs = [
        ("PL302", 14, 5, 25, 14),
        ("PL303", 14, 5, 25, 14),
        ("PL305", 8, 2, "max", 9),
        ("PL305", 6, 13, "max", 12),
        ("PL307", 14, 5, 25, 8),
        ("PL309", 14, 5, 25, 14),
        ("PL313", 14, 4, 10, 14),
        ("PL315", 14, 5, 25, 12),
        ("PL317", 15, 5, 25, 19),
        ("PL319", 14, 5, 25, 5),
    ]
    for sheet_name, row, col, rows, cols in specs:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        row_count = max(ws.max_row - row + 1, 1) if rows == "max" else int(rows)
        _clear_range(ws, row, col, row_count, cols)


def _write_product_values(wb, sheet_name: str, first_row: int, start_col: int, rows: list[list[Any]], metric_start_index: int) -> None:
    metrics = _product_report_metrics(rows, metric_start_index)
    _write_product_metrics(wb, sheet_name, first_row, start_col, metrics)


def _write_product_metrics(wb, sheet_name: str, first_row: int, start_col: int, metrics: dict[str, list[Any]]) -> None:
    if sheet_name not in wb.sheetnames:
        return
    order = _product_report_code_order()
    width = max((len(values) for values in metrics.values()), default=1)
    values = [_pad(metrics.get(code, []), width) for code in order]
    _write_matrix(wb[sheet_name], first_row, start_col, values)


def _product_report_metrics(rows: list[list[Any]], metric_start_index: int) -> dict[str, list[Any]]:
    out = {_text(row[0]): [_number(value) for value in row[metric_start_index:]] for row in rows}

    def sum_codes(codes: list[str]) -> list[float]:
        width = max((len(out.get(code, [])) for code in codes), default=0)
        values = [0.0] * width
        for code in codes:
            row = out.get(code, [])
            for i in range(width):
                values[i] += _number(row[i] if i < len(row) else 0)
        return values

    for prefix in ["A", "B", "C"]:
        out[prefix + "23"] = sum_codes([prefix + "21", prefix + "22"])
        out[prefix + "6"] = sum_codes([prefix + "1", prefix + "23", prefix + "3", prefix + "4", prefix + "5"])
    out["D"] = sum_codes(["A6", "B6", "C6"])
    return out


def _write_pl317(wb, rows: list[list[Any]]) -> None:
    metrics = _product_report_metrics(rows, 4)
    for values in metrics.values():
        while len(values) <= 18:
            values.append(0)
        values[18] = _number(values[15]) / _number(values[17]) * 100 if _number(values[17]) else 0
    _write_product_metrics(wb, "PL317", 15, 5, metrics)


def _write_pl319(wb, rows: list[list[Any]]) -> None:
    metrics = {_text(row[0]): [_number(value) for value in row[4:]] for row in rows}

    def aggregate(codes: list[str]) -> list[float]:
        balance = sum(_number(metrics.get(code, [0])[0]) for code in codes)
        interest = sum(_number((metrics.get(code, [0, 0]) + [0, 0])[1]) for code in codes)
        max_rates = [_number((metrics.get(code, [0, 0, 0, 0]) + [0, 0, 0, 0])[3]) for code in codes]
        min_rates = [_number((metrics.get(code, [0, 0, 0, 0, 0]) + [0, 0, 0, 0, 0])[4]) for code in codes]
        max_rates = [value for value in max_rates if value > 0]
        min_rates = [value for value in min_rates if value > 0]
        return [balance, interest, interest * 12 / balance * 100 if balance else 0, max(max_rates) if max_rates else 0, min(min_rates) if min_rates else 0]

    for prefix in ["A", "B", "C"]:
        metrics[prefix + "23"] = aggregate([prefix + "21", prefix + "22"])
        metrics[prefix + "6"] = aggregate([prefix + "1", prefix + "23", prefix + "3", prefix + "4", prefix + "5"])
    metrics["D"] = aggregate(["A6", "B6", "C6"])
    _write_product_metrics(wb, "PL319", 14, 5, metrics)


def _write_real_estate_values(wb, sheet_name: str, rows: list[list[Any]], metric_start_index: int) -> None:
    if sheet_name not in wb.sheetnames:
        return
    out = {_text(row[0]): [_number(value) for value in row[metric_start_index:]] for row in rows}

    def sum_codes(codes: list[str]) -> list[float]:
        width = max((len(out.get(code, [])) for code in codes), default=0)
        values = [0.0] * width
        for code in codes:
            row = out.get(code, [])
            for i in range(width):
                values[i] += _number(row[i] if i < len(row) else 0)
        return values

    out["C"] = sum_codes(["A", "B"])
    out["F"] = sum_codes(["D", "E"])
    out["I"] = sum_codes(["G", "H"])
    out["J"] = sum_codes(["C", "F", "I"])
    order = ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J"]
    width = max((len(values) for values in out.values()), default=1)
    _write_matrix(wb[sheet_name], 14, 4, [_pad(out.get(code, []), width) for code in order])


def _write_rows(wb, sheet_name: str, start_row: int, start_col: int, rows: list[list[Any]], width: int) -> None:
    if sheet_name not in wb.sheetnames:
        return
    if rows:
        _write_matrix(wb[sheet_name], start_row, start_col, [_pad(row, width) for row in rows])


def _write_matrix(ws, start_row: int, start_col: int, rows: list[list[Any]]) -> None:
    for r_idx, row in enumerate(rows, start_row):
        for c_offset, value in enumerate(row):
            cell = ws.cell(r_idx, start_col + c_offset)
            if isinstance(cell, MergedCell):
                continue
            cell.value = _excel_value(value)


def _clear_range(ws, start_row: int, start_col: int, rows: int, cols: int) -> None:
    for r_idx in range(start_row, start_row + max(rows, 0)):
        for c_idx in range(start_col, start_col + max(cols, 0)):
            cell = ws.cell(r_idx, c_idx)
            if isinstance(cell, MergedCell):
                continue
            cell.value = None


def _excel_value(value: Any) -> Any:
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _pad(values: list[Any], width: int) -> list[Any]:
    return values + [""] * max(width - len(values), 0)


def _product_report_code_order() -> list[str]:
    order = []
    for prefix in ["A", "B", "C"]:
        order.extend([prefix + "1", prefix + "21", prefix + "22", prefix + "23", prefix + "3", prefix + "4", prefix + "5", prefix + "6"])
    order.append("D")
    return order


def _product_rows() -> list[dict[str, Any]]:
    borrowers = [
        ("A", "개인", "개인대출"),
        ("B", "개인사업자", "개인 사업자 대출"),
        ("C", "법인", "법인대출"),
    ]
    products = [
        ("1", ["신용대출"], "신용대출", ""),
        ("21", ["부동산담보대출"], "부동산담보대출", "주택담보대출"),
        ("22", ["기타부동산 담보대출"], "부동산담보대출", "기타담보대출"),
        ("3", ["부동산 PF대출"], "부동산PF", ""),
        ("4", ["동산 담보대출", "기타담보대출", "임차보증금 담보대출", "NPL채권 담보대출", "ABL 대출", "유가증권 담보대출"], "기타담보(어음, 매출채권 제외)", ""),
        ("5", ["매출채권 담보대출"], "어음, 매출채권 담보", ""),
    ]
    return [
        {"code": prefix + code, "borrower": borrower, "borrowerLabel": borrower_label, "productLabel": product_label, "detailLabel": detail_label, "products": product_list}
        for prefix, borrower, borrower_label in borrowers
        for code, product_list, product_label, detail_label in products
    ]


def _real_estate_rows() -> list[dict[str, Any]]:
    return [
        {"code": "A", "borrower": "개인", "borrowerLabel": "개인 대출", "productLabel": "주택담보대출", "products": ["부동산담보대출"]},
        {"code": "B", "borrower": "개인", "borrowerLabel": "개인 대출", "productLabel": "기타담보대출", "products": ["기타부동산 담보대출"]},
        {"code": "D", "borrower": "개인사업자", "borrowerLabel": "개인 사업자 대출", "productLabel": "주택담보대출", "products": ["부동산담보대출"]},
        {"code": "E", "borrower": "개인사업자", "borrowerLabel": "개인 사업자 대출", "productLabel": "기타담보대출", "products": ["기타부동산 담보대출"]},
        {"code": "G", "borrower": "법인", "borrowerLabel": "법인 대출", "productLabel": "주택담보대출", "products": ["부동산담보대출"]},
        {"code": "H", "borrower": "법인", "borrowerLabel": "법인 대출", "productLabel": "기타담보대출", "products": ["기타부동산 담보대출"]},
    ]


def _investor_types() -> list[dict[str, str]]:
    return [
        {"key": "개인", "label": "일반투자자"},
        {"key": "소득적격", "label": "소득적격 개인투자자"},
        {"key": "전문", "label": "개인전문투자자"},
        {"key": "법인", "label": "일반법인"},
        {"key": "금융기관", "label": "금융기관"},
        {"key": "온투업자", "label": "온라인투자연계금융업자"},
    ]


def _invest_amount_bands() -> list[dict[str, float]]:
    return [
        {"label": "1백만원 이하", "min": float("-inf"), "max": 1_000_000},
        {"label": "1백만원 초과 5백만원 이하", "min": 1_000_000, "max": 5_000_000},
        {"label": "5백만원 초과 1천만원이하", "min": 5_000_000, "max": 10_000_000},
        {"label": "1천만원 초과 3천만원이하", "min": 10_000_000, "max": 30_000_000},
        {"label": "3천만원 초과 5천만원이하", "min": 30_000_000, "max": 50_000_000},
        {"label": "5천만원 초과", "min": 50_000_000, "max": float("inf")},
    ]


def _loan_amount_bands() -> list[dict[str, float | str]]:
    return [
        {"label": "5천만원 이하", "min": float("-inf"), "max": 50_000_000},
        {"label": "5천만원 초과 1억원이하", "min": 50_000_000, "max": 100_000_000},
        {"label": "1억원 초과 5억원 이하", "min": 100_000_000, "max": 500_000_000},
        {"label": "5억원 초과 30억원이하", "min": 500_000_000, "max": 3_000_000_000},
        {"label": "30억원 초과 70억원이하", "min": 3_000_000_000, "max": 7_000_000_000},
        {"label": "70억원 초과", "min": 7_000_000_000, "max": float("inf")},
    ]


def _ltv_bands() -> list[dict[str, float | str]]:
    return [
        {"label": "40% 이하", "min": float("-inf"), "max": 40},
        {"label": "40% 초과 50%이하", "min": 40, "max": 50},
        {"label": "50% 초과 60%이하", "min": 50, "max": 60},
        {"label": "60% 초과 70%이하", "min": 60, "max": 70},
        {"label": "70% 초과 100%이하", "min": 70, "max": 100},
        {"label": "100% 초과", "min": 100, "max": float("inf")},
    ]


def _term_bands() -> list[dict[str, float | str]]:
    return [
        {"label": "3개월 이하", "min": float("-inf"), "max": 3},
        {"label": "3개월 초과 6개월 이하", "min": 3, "max": 6},
        {"label": "6개월 초과 12개월 이하", "min": 6, "max": 12},
        {"label": "12개월 초과 24개월 이하", "min": 12, "max": 24},
        {"label": "24개월 초과", "min": 24, "max": float("inf")},
    ]


def _amount_band(amount: Any) -> str:
    value = _number(amount)
    hit = next((band for band in _loan_amount_bands() if value > band["min"] and value <= band["max"]), None)
    return str(hit["label"]) if hit else ""


def _ltv_band(ltv: Any) -> str:
    value = _number(ltv)
    hit = next((band for band in _ltv_bands() if value > band["min"] and value <= band["max"]), None)
    return str(hit["label"]) if hit else ""


def _term_band(term: Any) -> str:
    value = _number(term)
    hit = next((band for band in _term_bands() if value > band["min"] and value <= band["max"]), None)
    return str(hit["label"]) if hit else ""


def _overdue_band(row: dict[str, Any]) -> str:
    y = int(_number(row.get("기준년")))
    m = int(_number(row.get("기준월")))
    repay = _to_date(row.get("상환일"))
    if not y or not m or not repay:
        return ""
    end = _end_of_month(y, m)
    days = (end - repay).days
    if days <= 30:
        return ""
    months = _datedif_months(repay, end)
    if months < 3:
        return "3개월 이하"
    if months < 6:
        return "3개월 초과 6개월 이하"
    if months < 12:
        return "6개월 초과 12개월 이하"
    if months < 24:
        return "12개월 초과 24개월 이하"
    return "24개월 초과"


def _same_period(row: dict[str, Any], year: int | float, month: int | float) -> bool:
    return _number(row.get("기준년")) == _number(year) and _number(row.get("기준월")) == _number(month)


def _match_borrower(row: dict[str, Any], borrower: str) -> bool:
    return _text(row.get("대출자구분")).strip() == borrower


def _match_product(row: dict[str, Any], products: Iterable[str]) -> bool:
    raw = _text(row.get("세부상품") or row.get("상품분류_정리"))
    return raw in set(products)


def _normalize_investor_type(value: Any) -> str:
    raw = _text(value)
    if "소득" in raw:
        return "소득적격"
    if "전문" in raw:
        return "전문"
    if "법인" in raw:
        return "법인"
    if "금융" in raw or "여신" in raw:
        return "금융기관"
    if "온투" in raw:
        return "온투업자"
    return raw or "개인"


def _normalize_product(value: Any) -> str:
    raw = _text(value)
    if "PF" in raw:
        return "부동산 PF"
    if "기타부동산" in raw:
        return "기타부동산 담보대출"
    if "근저당" in raw:
        return "근저당질권"
    if "부동산담보" in raw or "부동산 담보" in raw:
        return "부동산담보대출"
    if "동산 담보" in raw:
        return "동산 담보대출"
    if "기타담보" in raw:
        return "기타담보대출"
    if "임차" in raw:
        return "임차보증금 담보대출"
    if "NPL" in raw:
        return "NPL채권 담보대출"
    if "ABL" in raw:
        return "ABL 대출"
    if "유가증권" in raw:
        return "유가증권 담보대출"
    if "매출채권" in raw or "어음" in raw:
        return "매출채권 담보대출"
    if "개인사업자" in raw and "신용" in raw:
        return "개인사업자 신용"
    if "법인" in raw and "신용" in raw:
        return "법인신용"
    if "신용" in raw:
        return "신용대출"
    return raw or "기타"


def _unique_count(rows: list[dict[str, Any]], key: str) -> int:
    return len({_text(row.get(key)) for row in rows if _text(row.get(key))})


def _unique_count_nonzero(rows: list[dict[str, Any]], key: str) -> int:
    return len(_unique_values(row.get(key) for row in rows))


def _unique_values(values: Iterable[Any]) -> list[str]:
    seen: dict[str, None] = {}
    for value in values:
        key = _text(value)
        if _valid_id(key):
            seen.setdefault(key, None)
    return list(seen.keys())


def _valid_id(value: Any) -> bool:
    compact = _text(value).replace(" ", "").replace("-", "")
    if not compact:
        return False
    try:
        return _number(compact) != 0
    except Exception:
        return True


def _sum(rows: Iterable[dict[str, Any]], key: str) -> float:
    return sum(_number(row.get(key)) for row in rows)


def _group_sum(rows: Iterable[dict[str, Any]], id_key: str, amount_key: str) -> dict[str, float]:
    out: dict[str, float] = {}
    for row in rows:
        key = _text(row.get(id_key))
        if not _valid_id(key):
            continue
        out[key] = out.get(key, 0) + _number(row.get(amount_key))
    return out


def _group_by(rows: Iterable[dict[str, Any]], key_fn) -> dict[str, list[dict[str, Any]]]:
    out: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        out[key_fn(row)].append(row)
    return dict(out)


def _sum_every_other(values: list[Any], start_index: int) -> float:
    return sum(_number(values[i]) for i in range(start_index, len(values), 2))


def _max(values: Iterable[Any]) -> float:
    nums = [_number(value) for value in values]
    nums = [value for value in nums if isfinite(value)]
    return max(nums) if nums else 0


def _max_contract_for_value(rows: list[dict[str, Any]], value_key: str, target_value: Any) -> float:
    target = _number(target_value)
    return _max(row.get("계약No") for row in rows if _number(row.get(value_key)) == target)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _value_or(value: Any, fallback: Any) -> Any:
    return fallback if value is None or value == "" else value


def _number(value: Any) -> float:
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return 0


def _period(year: int | float, month: int | float) -> dict[str, int]:
    return {"year": int(_number(year)), "month": int(_number(month))}


def _add_period(period: dict[str, int], delta_months: int) -> dict[str, int]:
    zero_based = period["year"] * 12 + (period["month"] - 1) + delta_months
    return {"year": zero_based // 12, "month": zero_based % 12 + 1}


def _end_of_month(year: int, month: int) -> date:
    if month == 12:
        return date(year, 12, 31)
    return date(year, month + 1, 1) - timedelta(days=1)


def _datedif_months(start_date: date, end_date: date) -> int:
    months = (end_date.year - start_date.year) * 12 + end_date.month - start_date.month
    if end_date.day < start_date.day:
        months -= 1
    return max(months, 0)


def _to_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return date(1899, 12, 30) + timedelta(days=int(round(value)))
    parsed = pd.to_datetime(value, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.date()


def _clean_product_name(name: Any) -> str:
    return _text(name).replace("[", "").replace("]", "")


def _mask_name(name: Any) -> str:
    raw = _text(name)
    return raw if len(raw) <= 1 else raw[0] + "*" * (len(raw) - 1)


def _business_no(value: Any) -> str:
    digits = "".join(ch for ch in _text(value) if ch.isdigit())
    return digits[:6] + "-" + digits[6:13] if len(digits) >= 13 else digits


def _yyyymmdd(value: Any) -> str:
    parsed = _to_date(value)
    if parsed:
        return parsed.strftime("%Y%m%d")
    return "".join(ch for ch in _text(value) if ch.isdigit())[:8]
